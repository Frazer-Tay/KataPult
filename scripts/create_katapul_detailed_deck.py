from pathlib import Path
import shutil

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches

from create_top3_feature_teaser import (
    W,
    H,
    C,
    F,
    background,
    browser,
    metric_card,
    module_tile,
    paragraph,
    rounded,
    scale_card,
    text,
)

ROOT = Path(__file__).resolve().parents[1]
PACKAGE = ROOT / "product-showcase-package"
SLIDE_DIR = PACKAGE / "katapult_detailed_deck_slides"
PPTX_PATH = PACKAGE / "KataPult_Detailed_Product_Showcase_8slides.pptx"


def save_slide(img, name):
    SLIDE_DIR.mkdir(parents=True, exist_ok=True)
    path = SLIDE_DIR / f"{name}.png"
    img.convert("RGB").save(path, quality=95)
    return path


def small_label(draw, x, y, label, active=False):
    fill = C["lime"] if active else "#172033"
    fg = C["bg"] if active else C["white"]
    rounded(draw, (x, y, x + 128, y + 34), fill, outline=None if active else C["line"], radius=11)
    text(draw, (x + 64, y + 17), label, F["tiny"], fg, anchor="mm")


def slide_title(draw, eyebrow, title, subtitle=None):
    rounded(draw, (72, 56, 184, 90), C["lime"], radius=10)
    text(draw, (128, 73), eyebrow.upper(), F["tiny"], C["bg"], anchor="mm")
    y = paragraph(draw, (72, 126), title, F["display"], C["white"], 560, gap=4)
    if subtitle:
        paragraph(draw, (76, y + 20), subtitle, F["body"], C["cyan"], 560, gap=5)


def slide_cover():
    img = background()
    d = ImageDraw.Draw(img, "RGBA")
    paragraph(d, (72, 84), "KataPult", F["display"], C["white"], 560)
    paragraph(d, (76, 166), "Bahasa I2 practice by component.", F["h2"], C["cyan"], 560)
    paragraph(d, (76, 226), "A structured way to choose weak areas, practise in short drills, and see usage.", F["body"], C["muted"], 440)
    browser(img, (560, 76, 1210, 570), "01-home.png", scale=0.96, glow=True)
    metric_card(d, (76, 390, 236, 466), "1,554*", "drill / daily pool")
    metric_card(d, (260, 390, 420, 466), "971**", "custom-test bank")
    metric_card(d, (76, 490, 236, 566), "40++", "learners")
    metric_card(d, (260, 490, 420, 566), "4+", "entities")
    text(d, (76, 646), "From scattered prep to repeatable practice.", F["body_b"], C["white"])
    text(d, (610, 646), "* Vocab + imbuhan + persamaan. ** Imbuhan + persamaan only.", F["tiny"], C["muted"])
    return img


def slide_problem():
    img = background()
    d = ImageDraw.Draw(img, "RGBA")
    slide_title(d, "Problem", "Content existed.\nPractice was fragmented.", "The hard part was not finding materials. It was turning them into repeatable exam practice.")
    cards = [
        ("Notes", "Personal summaries"),
        ("PDFs", "Reference files"),
        ("Screenshots", "Saved examples"),
        ("Links", "Useful pages"),
        ("Quizlet", "Existing sets"),
        ("LLM prompts", "Extra practice"),
    ]
    positions = [(660, 116), (910, 116), (660, 226), (910, 226), (660, 336), (910, 336)]
    for (label, detail), (x, y) in zip(cards, positions):
        module_tile(img, (x, y, x + 220, y + 72), label, detail, active=False)
    rounded(d, (660, 500, 1164, 598), "#101827", outline=C["line"], radius=22)
    text(d, (690, 528), "The gap was practice, not content.", F["h2"], C["white"])
    text(d, (690, 562), "Reading across resources was not the same as retrieval and feedback.", F["small"], C["muted"])
    return img


def slide_workflow():
    img = background()
    d = ImageDraw.Draw(img, "RGBA")
    slide_title(d, "Solution", "A repeatable practice loop.", "KataPult turns preparation into a simple routine: choose, drill, review, track.")
    browser(img, (650, 98, 1210, 570), "01-home.png", scale=0.98, glow=True)
    steps = [
        ("1", "Choose component", "Start with the exact skill to practise."),
        ("2", "Drill actively", "Answer questions instead of rereading."),
        ("3", "Get feedback", "Know what to repeat next."),
        ("4", "Track practice", "See section visits and time spent."),
    ]
    for i, (num, title, detail) in enumerate(steps):
        y = 318 + i * 78
        rounded(d, (86, y, 126, y + 40), C["lime"], radius=10)
        text(d, (106, y + 20), num, F["small_b"], C["bg"], anchor="mm")
        text(d, (150, y + 1), title, F["body_b"], C["white"])
        text(d, (150, y + 30), detail, F["small"], C["muted"])
    return img


def slide_scale():
    img = background()
    d = ImageDraw.Draw(img, "RGBA")
    slide_title(d, "Feature 1", "Exam-skill modules\nwith visible scale.", "Drill items and revision modules are counted separately.")
    browser(img, (660, 94, 1210, 498), "04-persamaan.png", scale=0.98, glow=True)
    counts = [
        ("583", "Vocabulary items"),
        ("666", "Imbuhan items"),
        ("305", "Persamaan items"),
        ("297", "Karangan phrases"),
        ("8", "Surat rasmi sections"),
        ("23", "Flashcard sets"),
    ]
    for i, (value, label) in enumerate(counts):
        x = 76 + (i % 3) * 188
        y = 380 + (i // 3) * 102
        scale_card(d, (x, y, x + 164, y + 78), value, label, active=label.startswith("Imbuhan") or label.startswith("Persamaan"))
    text(d, (662, 548), "Math: 583 + 666 + 305 = 1,554 drill/daily pool.", F["body_b"], C["white"])
    paragraph(d, (662, 582), "Karangan, surat rasmi and flashcards are supporting revision modules, not part of the drill-question pool.", F["small"], C["muted"], 500, gap=3)
    return img


def slide_custom_drills():
    img = background()
    d = ImageDraw.Draw(img, "RGBA")
    slide_title(d, "Feature 2", "Custom drills for weak areas.", "The 971-item custom-test bank is a subset: 666 imbuhan + 305 persamaan.")
    browser(img, (620, 92, 1210, 596), "06-test-setup.png", scale=0.98, glow=True)
    metric_card(d, (76, 356, 250, 436), "971**", "custom-test bank")
    metric_card(d, (276, 356, 450, 436), "5/10/20", "drill lengths")
    steps = [
        ("Choose component", "Imbuhan or persamaan"),
        ("Set question count", "Short session by design"),
        ("Practise actively", "Answer first, then review"),
    ]
    for i, (title, detail) in enumerate(steps):
        y = 474 + i * 62
        module_tile(img, (76, y, 450, y + 52), title, detail, active=i == 1)
    return img


def slide_habit_tracking():
    img = background()
    d = ImageDraw.Draw(img, "RGBA")
    slide_title(d, "Feature 3", "Daily habit plus tracking.", "Daily challenge draws from the 1,554-item pool: vocabulary + imbuhan + persamaan.")
    browser(img, (620, 86, 1210, 590), "07-daily-challenge.png", scale=0.98, glow=True)
    metric_card(d, (76, 354, 250, 434), "10", "daily questions")
    metric_card(d, (276, 354, 450, 434), "1,554*", "daily challenge pool")
    module_tile(img, (76, 478, 450, 534), "Habit cues", "Streak, XP, level and lives", active=True)
    module_tile(img, (76, 550, 450, 606), "Usage visibility", "Section visits and time spent", active=False)
    text(d, (76, 642), "The point is not long study blocks. It is repeated, visible practice.", F["body_b"], C["white"])
    return img


def slide_demo_flow():
    img = background()
    d = ImageDraw.Draw(img, "RGBA")
    rounded(d, (72, 56, 156, 90), C["lime"], radius=10)
    text(d, (114, 73), "DEMO", F["tiny"], C["bg"], anchor="mm")
    paragraph(d, (72, 126), "10-minute learner journey.", F["display"], C["white"], 520, gap=4)
    paragraph(d, (76, 210), "Use this flow when walking someone through the app.", F["body"], C["cyan"], 520, gap=5)
    shots = [
        ("Open dashboard", "01-home.png"),
        ("Practise vocabulary", "02-vocabulary.png"),
        ("Drill imbuhan", "03-imbuhan.png"),
        ("Run custom test", "06-test-setup.png"),
    ]
    for i, (label, src) in enumerate(shots):
        x = 76 + (i % 2) * 550
        y = 330 + (i // 2) * 170
        browser(img, (x, y, x + 470, y + 132), src, scale=1.0, glow=False)
        rounded(d, (x, y - 44, x + 42, y - 10), C["lime"], radius=9)
        text(d, (x + 21, y - 27), str(i + 1), F["tiny"], C["bg"], anchor="mm")
        text(d, (x + 58, y - 39), label, F["body_b"], C["white"])
    return img


def slide_value():
    img = background()
    d = ImageDraw.Draw(img, "RGBA")
    slide_title(d, "Value", "A reusable practice surface.", "KataPult turns individual prep material into a shared platform with visible usage.")
    browser(img, (632, 96, 1210, 500), "01-home.png", scale=0.98, glow=True)
    metric_card(d, (76, 348, 250, 428), "40++", "learners reached")
    metric_card(d, (276, 348, 450, 428), "4+", "entities")
    metric_card(d, (76, 456, 250, 536), "1,554*", "drill / daily pool")
    metric_card(d, (276, 456, 450, 536), "Ready", "usage analytics")
    rounded(d, (632, 540, 1186, 620), "#101827", outline=C["line"], radius=20)
    text(d, (662, 562), "Before: scattered, individual, difficult to observe.", F["body_b"], C["white"])
    text(d, (662, 590), "After: structured, reusable, and measurable.", F["body_b"], C["lime"])
    text(d, (76, 650), "* Vocab + imbuhan + persamaan. ** Custom tests use imbuhan + persamaan only.", F["small"], C["muted"])
    return img


SLIDES = [
    ("01_cover", slide_cover),
    ("02_problem", slide_problem),
    ("03_workflow", slide_workflow),
    ("04_scale", slide_scale),
    ("05_custom_drills", slide_custom_drills),
    ("06_habit_tracking", slide_habit_tracking),
    ("07_demo_flow", slide_demo_flow),
    ("08_value", slide_value),
]


def build_ppt(slide_paths):
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    for path in slide_paths:
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(path), 0, 0, width=prs.slide_width, height=prs.slide_height)
    prs.save(PPTX_PATH)


def main():
    if SLIDE_DIR.exists():
        shutil.rmtree(SLIDE_DIR)
    SLIDE_DIR.mkdir(parents=True, exist_ok=True)
    slide_paths = []
    for name, factory in SLIDES:
        slide_paths.append(save_slide(factory(), name))
    build_ppt(slide_paths)
    print(PPTX_PATH)
    print(SLIDE_DIR)
    print(f"slides={len(slide_paths)}")


if __name__ == "__main__":
    main()
