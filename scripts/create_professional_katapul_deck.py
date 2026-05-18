from pathlib import Path
from zipfile import ZipFile
import shutil

from PIL import Image, ImageDraw, ImageFont, ImageFilter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "award-demo-assets"
PACKAGE = ROOT / "product-showcase-package"
SLIDE_DIR = PACKAGE / "katapult_professional_deck_slides"
PPTX_PATH = PACKAGE / "KataPult_Professional_Product_Deck_8slides.pptx"

W, H = 1600, 900

COL = {
    "bg": "#F7F9FC",
    "paper": "#FFFFFF",
    "navy": "#111827",
    "ink": "#172033",
    "muted": "#647084",
    "line": "#D9E0EA",
    "soft": "#EEF3F8",
    "lime": "#A6D83A",
    "cyan": "#36BDD3",
    "purple": "#6D5BD0",
    "amber": "#F5B84B",
}


def font(size, weight="regular"):
    paths = {
        "light": ["C:/Windows/Fonts/segoeuisl.ttf", "C:/Windows/Fonts/segoeuil.ttf", "C:/Windows/Fonts/calibril.ttf"],
        "regular": ["C:/Windows/Fonts/segoeui.ttf", "C:/Windows/Fonts/calibri.ttf", "C:/Windows/Fonts/arial.ttf"],
        "bold": ["C:/Windows/Fonts/segoeuib.ttf", "C:/Windows/Fonts/calibrib.ttf", "C:/Windows/Fonts/arialbd.ttf"],
    }
    for p in paths[weight]:
        if Path(p).exists():
            return ImageFont.truetype(p, size)
    return ImageFont.load_default()


F = {
    "kicker": font(17, "bold"),
    "display": font(56, "light"),
    "h1": font(44, "light"),
    "h2": font(28, "bold"),
    "body": font(22, "regular"),
    "body_b": font(22, "bold"),
    "small": font(17, "regular"),
    "small_b": font(17, "bold"),
    "tiny": font(13, "regular"),
    "metric": font(42, "bold"),
}


def rgb(hex_color):
    hex_color = hex_color.lstrip("#")
    return tuple(int(hex_color[i : i + 2], 16) for i in (0, 2, 4))


def rgba(hex_color, alpha=255):
    return rgb(hex_color) + (alpha,)


def rounded(draw, box, fill, outline=None, radius=22, width=1):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def text(draw, xy, value, fnt, fill, anchor=None):
    draw.text(xy, value, font=fnt, fill=fill, anchor=anchor)


def wrap_lines(draw, value, fnt, max_width):
    lines, cur = [], ""
    for word in str(value).split():
        candidate = (cur + " " + word).strip()
        if draw.textbbox((0, 0), candidate, font=fnt)[2] <= max_width:
            cur = candidate
        else:
            if cur:
                lines.append(cur)
            cur = word
    if cur:
        lines.append(cur)
    return lines


def para(draw, xy, value, fnt, fill, max_width, gap=7):
    x, y = xy
    for line in wrap_lines(draw, value, fnt, max_width):
        text(draw, (x, y), line, fnt, fill)
        y += fnt.size + gap
    return y


def shadow(layer, box, blur=18, alpha=45, radius=24):
    sh = Image.new("RGBA", layer.size, (0, 0, 0, 0))
    sd = ImageDraw.Draw(sh, "RGBA")
    sd.rounded_rectangle(box, radius=radius, fill=(15, 23, 42, alpha))
    sh = sh.filter(ImageFilter.GaussianBlur(blur))
    layer.alpha_composite(sh)


def base(slide_no, section):
    img = Image.new("RGBA", (W, H), rgba(COL["bg"]))
    d = ImageDraw.Draw(img, "RGBA")
    d.rectangle((0, 0, W, 88), fill=COL["navy"])
    d.rectangle((0, 88, W, 92), fill=COL["lime"])
    text(d, (76, 30), "KataPult", F["body_b"], COL["paper"])
    text(d, (204, 32), section, F["small"], "#B8C2D2")
    text(d, (W - 76, 32), f"{slide_no:02d}", F["small_b"], "#B8C2D2", anchor="ra")
    return img


def title_block(draw, kicker, title, subtitle=None, x=76, y=128, width=660):
    text(draw, (x, y), kicker.upper(), F["kicker"], COL["cyan"])
    y = para(draw, (x, y + 34), title, F["display"], COL["ink"], width, gap=5)
    if subtitle:
        para(draw, (x, y + 18), subtitle, F["body"], COL["muted"], width, gap=6)


def card(draw, box, title=None, body=None, accent=None):
    rounded(draw, box, COL["paper"], outline=COL["line"], radius=22)
    if accent:
        draw.rounded_rectangle((box[0] + 22, box[1] + 22, box[0] + 28, box[3] - 22), radius=3, fill=accent)
        tx = box[0] + 48
    else:
        tx = box[0] + 26
    if title:
        text(draw, (tx, box[1] + 22), title, F["body_b"], COL["ink"])
    if body:
        para(draw, (tx, box[1] + 58), body, F["small"], COL["muted"], box[2] - tx - 24, gap=4)


def metric(draw, box, value, label, note=None, accent=COL["lime"]):
    rounded(draw, box, COL["paper"], outline=COL["line"], radius=20)
    text(draw, (box[0] + 26, box[1] + 22), value, F["metric"], accent)
    text(draw, (box[0] + 26, box[1] + 72), label, F["small_b"], COL["ink"])
    if note:
        para(draw, (box[0] + 26, box[1] + 102), note, F["tiny"], COL["muted"], box[2] - box[0] - 52, gap=3)


def screenshot(name, size):
    im = Image.open(ASSETS / name).convert("RGBA")
    im.thumbnail(size)
    canvas = Image.new("RGBA", size, rgba(COL["paper"]))
    canvas.alpha_composite(im, ((size[0] - im.width) // 2, (size[1] - im.height) // 2))
    return canvas


def browser(layer, box, name):
    shadow(layer, box, blur=18, alpha=52, radius=26)
    d = ImageDraw.Draw(layer, "RGBA")
    rounded(d, box, COL["paper"], outline=COL["line"], radius=26)
    d.rounded_rectangle((box[0], box[1], box[2], box[1] + 44), radius=26, fill=COL["soft"])
    for i, c in enumerate(["#F16B6B", "#F5C445", "#4FD08B"]):
        d.ellipse((box[0] + 24 + i * 26, box[1] + 16, box[0] + 38 + i * 26, box[1] + 30), fill=c)
    d.rounded_rectangle((box[0] + 160, box[1] + 16, box[2] - 34, box[1] + 30), radius=7, fill="#DCE5EF")
    shot = screenshot(name, (box[2] - box[0] - 34, box[3] - box[1] - 58))
    layer.alpha_composite(shot, (box[0] + 17, box[1] + 52))


def formula(draw, box, title, body):
    rounded(draw, box, "#F2F7FB", outline=COL["line"], radius=18)
    text(draw, (box[0] + 24, box[1] + 20), title, F["small_b"], COL["ink"])
    para(draw, (box[0] + 24, box[1] + 52), body, F["tiny"], COL["muted"], box[2] - box[0] - 48, gap=4)


def slide_01():
    img = base(1, "Executive Summary")
    d = ImageDraw.Draw(img, "RGBA")
    title_block(d, "Product overview", "Bahasa I2 practice,\nstructured by component.", "KataPult helps learners move from scattered resources to targeted, repeatable practice.", width=690)
    browser(img, (870, 150, 1510, 620), "01-home.png")
    metric(d, (76, 520, 286, 675), "1,554*", "drill / daily pool", "Vocabulary + imbuhan + persamaan")
    metric(d, (316, 520, 526, 675), "971**", "custom-test bank", "Imbuhan + persamaan only")
    metric(d, (556, 520, 766, 675), "40++", "learners", "Across at least 4 entities")
    text(d, (76, 806), "* Drill pool excludes revision-only modules. ** Custom tests draw from question-form components.", F["tiny"], COL["muted"])
    return img


def slide_02():
    img = base(2, "Problem")
    d = ImageDraw.Draw(img, "RGBA")
    title_block(d, "Context", "The issue was not content.\nIt was practice quality.", "Useful material existed, but preparation was spread across notes, PDFs, screenshots, links, Quizlet sets and LLM prompts.", width=740)
    cols = [
        ("Scattered inputs", "Resources lived in separate places, making it hard to restart quickly."),
        ("Passive revision", "Reading and collecting examples did not force retrieval."),
        ("No shared visibility", "There was no easy way to see what was practised or where time was spent."),
    ]
    for i, (h, b) in enumerate(cols):
        x = 76 + i * 485
        card(d, (x, 545, x + 430, 725), h, b, accent=[COL["cyan"], COL["lime"], COL["amber"]][i])
    formula(d, (940, 190, 1488, 386), "Operational problem", "Each learner could prepare independently, but the organisation did not get a reusable practice surface or usage signal from that effort.")
    return img


def slide_03():
    img = base(3, "Solution Model")
    d = ImageDraw.Draw(img, "RGBA")
    title_block(d, "Workflow", "One repeatable practice loop.", "The app is organised around what a learner does in a short session.", width=640)
    steps = [
        ("1", "Choose component", "Vocabulary, imbuhan, persamaan, writing, surat rasmi or flashcards."),
        ("2", "Practise actively", "Answer targeted questions instead of rereading long notes."),
        ("3", "Review feedback", "See what to repeat and which weak area to revisit."),
        ("4", "Track usage", "Capture section visits and time spent as adoption evidence."),
    ]
    for i, (num, h, b) in enumerate(steps):
        x = 76 + (i % 2) * 390
        y = 420 + (i // 2) * 155
        rounded(d, (x, y, x + 54, y + 54), COL["lime"], radius=14)
        text(d, (x + 27, y + 27), num, F["small_b"], COL["navy"], anchor="mm")
        card(d, (x + 74, y - 8, x + 370, y + 120), h, b)
    browser(img, (910, 164, 1510, 650), "01-home.png")
    return img


def slide_04():
    img = base(4, "Practice Bank")
    d = ImageDraw.Draw(img, "RGBA")
    title_block(d, "Scale and definitions", "The numbers are counted\nby use case.", "This slide prevents the common confusion between total drill pool, custom-test bank and revision modules.", width=760)
    metric(d, (76, 395, 316, 555), "1,554*", "drill / daily pool", "583 vocabulary + 666 imbuhan + 305 persamaan")
    metric(d, (346, 395, 586, 555), "971**", "custom-test bank", "666 imbuhan + 305 persamaan")
    metric(d, (616, 395, 856, 555), "3", "revision surfaces", "Karangan phrases, surat rasmi and flashcards", accent=COL["cyan"])
    table_x, table_y = 940, 150
    rows = [
        ("Vocabulary", "583", "Drill / daily pool"),
        ("Imbuhan", "666", "Drill / daily + custom tests"),
        ("Persamaan", "305", "Drill / daily + custom tests"),
        ("Karangan", "297", "Writing phrases / revision"),
        ("Surat rasmi", "8 sections", "Templates and structure"),
        ("Flashcards", "23 sets", "Recall and review"),
    ]
    rounded(d, (table_x, table_y, 1500, 680), COL["paper"], outline=COL["line"], radius=22)
    text(d, (table_x + 28, table_y + 26), "Component", F["small_b"], COL["muted"])
    text(d, (table_x + 248, table_y + 26), "Scale", F["small_b"], COL["muted"])
    text(d, (table_x + 374, table_y + 26), "Use", F["small_b"], COL["muted"])
    for i, row in enumerate(rows):
        y = table_y + 72 + i * 62
        d.line((table_x + 24, y - 16, 1476, y - 16), fill=COL["line"], width=1)
        text(d, (table_x + 28, y), row[0], F["small_b"], COL["ink"])
        text(d, (table_x + 248, y), row[1], F["small_b"], COL["ink"])
        text(d, (table_x + 374, y), row[2], F["small"], COL["muted"])
    text(d, (76, 622), "* 583 + 666 + 305 = 1,554. ** 666 + 305 = 971.", F["body_b"], COL["ink"])
    para(d, (76, 662), "Karangan, surat rasmi and flashcards are shown separately because they are revision surfaces, not custom-test question banks.", F["small"], COL["muted"], 770, gap=5)
    return img


def slide_05():
    img = base(5, "Feature 1")
    d = ImageDraw.Draw(img, "RGBA")
    title_block(d, "Exam-skill modules", "Learners practise\nthe exact weak area.", "This is not a generic language app: each section maps to an I2-relevant practice need.", width=660)
    browser(img, (890, 144, 1510, 610), "04-persamaan.png")
    modules = [
        ("Vocabulary", "Meaning and usage", "583 items"),
        ("Imbuhan", "Word formation", "666 items"),
        ("Persamaan", "Synonym precision", "305 items"),
        ("Karangan", "Ready phrases", "297 phrases"),
        ("Surat rasmi", "Structure and tone", "8 sections"),
        ("Flashcards", "Recall and review", "23 sets"),
    ]
    for i, (h, b, m) in enumerate(modules):
        x = 76 + (i % 2) * 345
        y = 392 + (i // 2) * 120
        card(d, (x, y, x + 305, y + 94), h, f"{b}\n{m}", accent=COL["lime"] if h in {"Imbuhan", "Persamaan"} else COL["cyan"])
    return img


def slide_06():
    img = base(6, "Feature 2")
    d = ImageDraw.Draw(img, "RGBA")
    title_block(d, "Focused custom drills", "A weak area becomes\na short test.", "Custom tests use the 971-item bank from imbuhan and persamaan question items.", width=700)
    browser(img, (840, 132, 1510, 626), "06-test-setup.png")
    formula(d, (76, 392, 740, 484), "Custom-test math", "666 imbuhan items + 305 persamaan items = 971 custom-test items.")
    steps = [
        ("Choose", "Select imbuhan or persamaan."),
        ("Set length", "Run 5, 10 or 20 questions depending on available time."),
        ("Review", "Use mistakes to decide what to repeat next."),
    ]
    for i, (h, b) in enumerate(steps):
        x = 76 + i * 255
        card(d, (x, 540, x + 220, 690), h, b, accent=COL["lime"] if i == 1 else COL["cyan"])
    return img


def slide_07():
    img = base(7, "Feature 3")
    d = ImageDraw.Draw(img, "RGBA")
    title_block(d, "Daily practice and usage", "Keep sessions small.\nMake behaviour visible.", "Daily challenge draws from the 1,554-item pool; tracking records what users actually visit and practise.", width=760)
    browser(img, (890, 132, 1510, 626), "07-daily-challenge.png")
    metric(d, (76, 414, 286, 564), "10", "daily questions", "Small enough for a quick session")
    metric(d, (316, 414, 526, 564), "1,554*", "daily pool", "Vocabulary + imbuhan + persamaan")
    metric(d, (556, 414, 766, 564), "5-10", "minutes", "Designed around pockets of time")
    card(d, (76, 622, 766, 746), "What can be observed", "Section visits, time spent by section, total session duration, active days and adoption across entities.", accent=COL["cyan"])
    return img


def slide_08():
    img = base(8, "Demo and Value")
    d = ImageDraw.Draw(img, "RGBA")
    title_block(d, "How to present it", "Demo the learner flow,\nthen show the evidence.", "A clean 10-minute prep path, followed by the usage signal.", width=760)
    flow = [
        ("1", "Open dashboard", "Shows all practice paths in one place."),
        ("2", "Pick weak area", "Jump straight to vocabulary, imbuhan or persamaan."),
        ("3", "Run custom drill", "Turn a weak area into a short test."),
        ("4", "Check usage", "Show section and time tracking once enough data accumulates."),
    ]
    for i, (num, h, b) in enumerate(flow):
        y = 388 + i * 88
        rounded(d, (86, y, 126, y + 44), COL["lime"], radius=12)
        text(d, (106, y + 22), num, F["small_b"], COL["navy"], anchor="mm")
        text(d, (152, y), h, F["body_b"], COL["ink"])
        text(d, (152, y + 30), b, F["small"], COL["muted"])
    browser(img, (890, 132, 1510, 516), "01-home.png")
    metric(d, (890, 570, 1060, 716), "40++", "learners", "Known adoption signal")
    metric(d, (1085, 570, 1255, 716), "4+", "entities", "Cross-team use")
    metric(d, (1280, 570, 1510, 716), "Ready", "analytics", "Usage evidence as logging matures")
    return img


SLIDES = [
    ("01_executive_summary", slide_01),
    ("02_problem", slide_02),
    ("03_solution_loop", slide_03),
    ("04_practice_bank_math", slide_04),
    ("05_exam_modules", slide_05),
    ("06_custom_drills", slide_06),
    ("07_daily_usage", slide_07),
    ("08_demo_value", slide_08),
]


def save_slide(img, name):
    SLIDE_DIR.mkdir(parents=True, exist_ok=True)
    out = SLIDE_DIR / f"{name}.png"
    img.convert("RGB").save(out, quality=95)
    return out


def build_ppt(paths):
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    def sx(px):
        return Inches(px / W * 13.333333)

    def sy(px):
        return Inches(px / H * 7.5)

    def pbox(box):
        x1, y1, x2, y2 = box
        return sx(x1), sy(y1), sx(x2 - x1), sy(y2 - y1)

    def ppt_color(hex_color):
        return RGBColor(*rgb(hex_color))

    def style_shape(shape, fill, line=None, width=1):
        shape.fill.solid()
        shape.fill.fore_color.rgb = ppt_color(fill)
        if line:
            shape.line.color.rgb = ppt_color(line)
            shape.line.width = Pt(width)
        else:
            shape.line.fill.background()
        return shape

    def rect(slide, box, fill, line=None, rounded=False, width=1):
        shp = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
            *pbox(box),
        )
        return style_shape(shp, fill, line, width)

    def oval(slide, box, fill):
        shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, *pbox(box))
        return style_shape(shp, fill)

    def add_text(slide, box, value, size, color=COL["ink"], bold=False, font_name="Aptos", align=None):
        tx = slide.shapes.add_textbox(*pbox(box))
        tf = tx.text_frame
        tf.clear()
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = 0
        tf.margin_right = 0
        tf.margin_top = 0
        tf.margin_bottom = 0
        tf.text = value
        for paragraph in tf.paragraphs:
            paragraph.space_after = Pt(0)
            paragraph.line_spacing = 1.05
            if align is not None:
                paragraph.alignment = align
            for run in paragraph.runs:
                run.font.name = font_name
                run.font.size = Pt(size)
                run.font.bold = bold
                run.font.color.rgb = ppt_color(color)
        return tx

    def base_native(slide_no, section):
        slide = prs.slides.add_slide(blank)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = ppt_color(COL["bg"])
        rect(slide, (0, 0, W, 88), COL["navy"])
        rect(slide, (0, 88, W, 92), COL["lime"])
        add_text(slide, (76, 30, 190, 62), "KataPult", 12, COL["paper"], True, "Aptos")
        add_text(slide, (204, 32, 560, 62), section, 10.5, "#B8C2D2", False, "Aptos")
        add_text(slide, (1450, 32, 1524, 62), f"{slide_no:02d}", 10.5, "#B8C2D2", True, "Aptos", PP_ALIGN.RIGHT)
        return slide

    def title_native(slide, kicker, title, subtitle=None, x=76, y=128, width=660):
        add_text(slide, (x, y, x + width, y + 24), kicker.upper(), 10, COL["cyan"], True, "Aptos")
        lines = title.count("\n") + 1
        title_h = 66 * lines
        add_text(slide, (x, y + 42, x + width, y + 42 + title_h), title, 34, COL["ink"], False, "Aptos Display")
        if subtitle:
            add_text(slide, (x, y + 54 + title_h, x + width, y + 54 + title_h + 82), subtitle, 15, COL["muted"], False, "Aptos")

    def card_native(slide, box, title=None, body=None, accent=None):
        rect(slide, box, COL["paper"], COL["line"], rounded=True, width=0.7)
        tx = box[0] + 26
        if accent:
            rect(slide, (box[0] + 22, box[1] + 22, box[0] + 29, box[3] - 22), accent, rounded=True)
            tx = box[0] + 48
        if title:
            add_text(slide, (tx, box[1] + 22, box[2] - 22, box[1] + 50), title, 12.5, COL["ink"], True, "Aptos")
        if body:
            add_text(slide, (tx, box[1] + 60, box[2] - 24, box[3] - 18), body, 10.5, COL["muted"], False, "Aptos")

    def metric_native(slide, box, value, label, note=None, accent=COL["lime"]):
        rect(slide, box, COL["paper"], COL["line"], rounded=True, width=0.7)
        value_size = 26 if len(value) <= 6 else 23
        add_text(slide, (box[0] + 26, box[1] + 22, box[2] - 24, box[1] + 72), value, value_size, accent, True, "Aptos Display")
        add_text(slide, (box[0] + 26, box[1] + 78, box[2] - 24, box[1] + 104), label, 10.5, COL["ink"], True, "Aptos")
        if note:
            add_text(slide, (box[0] + 26, box[1] + 110, box[2] - 24, box[3] - 12), note, 8.2, COL["muted"], False, "Aptos")

    def formula_native(slide, box, title, body):
        rect(slide, box, "#F2F7FB", COL["line"], rounded=True, width=0.7)
        add_text(slide, (box[0] + 24, box[1] + 20, box[2] - 24, box[1] + 45), title, 10.5, COL["ink"], True, "Aptos")
        add_text(slide, (box[0] + 24, box[1] + 56, box[2] - 24, box[3] - 16), body, 8.6, COL["muted"], False, "Aptos")

    def add_picture_fit(slide, path, box):
        im = Image.open(path)
        max_w = box[2] - box[0]
        max_h = box[3] - box[1]
        ratio = min(max_w / im.width, max_h / im.height)
        pic_w = im.width * ratio
        pic_h = im.height * ratio
        left = box[0] + (max_w - pic_w) / 2
        top = box[1] + (max_h - pic_h) / 2
        slide.shapes.add_picture(str(path), sx(left), sy(top), width=sx(pic_w), height=sy(pic_h))

    def browser_native(slide, box, name):
        rect(slide, (box[0] + 9, box[1] + 13, box[2] + 9, box[3] + 13), "#DCE4EE", rounded=True)
        rect(slide, box, COL["paper"], COL["line"], rounded=True, width=0.7)
        rect(slide, (box[0], box[1], box[2], box[1] + 44), COL["soft"], rounded=True)
        for i, c in enumerate(["#F16B6B", "#F5C445", "#4FD08B"]):
            oval(slide, (box[0] + 24 + i * 26, box[1] + 16, box[0] + 38 + i * 26, box[1] + 30), c)
        rect(slide, (box[0] + 160, box[1] + 16, box[2] - 34, box[1] + 30), "#DCE5EF", rounded=True)
        add_picture_fit(slide, ASSETS / name, (box[0] + 17, box[1] + 52, box[2] - 17, box[3] - 6))

    def build_slide_01():
        slide = base_native(1, "Executive Summary")
        title_native(slide, "Product overview", "Bahasa I2 practice,\nstructured by component.", "KataPult helps learners move from scattered resources to targeted, repeatable practice.", width=690)
        browser_native(slide, (870, 150, 1510, 620), "01-home.png")
        metric_native(slide, (76, 520, 286, 675), "1,554*", "drill / daily pool", "Vocabulary + imbuhan + persamaan")
        metric_native(slide, (316, 520, 526, 675), "971**", "custom-test bank", "Imbuhan + persamaan only")
        metric_native(slide, (556, 520, 766, 675), "40++", "learners", "Across at least 4 entities")
        add_text(slide, (76, 806, 980, 832), "* Drill pool excludes revision-only modules. ** Custom tests draw from question-form components.", 8.2, COL["muted"])

    def build_slide_02():
        slide = base_native(2, "Problem")
        title_native(slide, "Context", "The issue was not content.\nIt was practice quality.", "Useful material existed, but preparation was spread across notes, PDFs, screenshots, links, Quizlet sets and LLM prompts.", width=740)
        formula_native(slide, (940, 190, 1488, 386), "Operational problem", "Each learner could prepare independently, but the organisation did not get a reusable practice surface or usage signal from that effort.")
        cols = [
            ("Scattered inputs", "Resources lived in separate places, making it hard to restart quickly."),
            ("Passive revision", "Reading and collecting examples did not force retrieval."),
            ("No shared visibility", "There was no easy way to see what was practised or where time was spent."),
        ]
        for i, (h, b) in enumerate(cols):
            x = 76 + i * 485
            card_native(slide, (x, 545, x + 430, 725), h, b, [COL["cyan"], COL["lime"], COL["amber"]][i])

    def build_slide_03():
        slide = base_native(3, "Solution Model")
        title_native(slide, "Workflow", "One repeatable practice loop.", "The app is organised around what a learner does in a short session.", width=640)
        steps = [
            ("1", "Choose component", "Vocabulary, imbuhan, persamaan, writing, surat rasmi or flashcards."),
            ("2", "Practise actively", "Answer targeted questions instead of rereading long notes."),
            ("3", "Review feedback", "See what to repeat and which weak area to revisit."),
            ("4", "Track usage", "Capture section visits and time spent as adoption evidence."),
        ]
        for i, (num, h, b) in enumerate(steps):
            x = 76 + (i % 2) * 390
            y = 420 + (i // 2) * 155
            rect(slide, (x, y, x + 54, y + 54), COL["lime"], rounded=True)
            add_text(slide, (x, y + 13, x + 54, y + 41), num, 10.5, COL["navy"], True, "Aptos", PP_ALIGN.CENTER)
            card_native(slide, (x + 74, y - 8, x + 370, y + 120), h, b)
        browser_native(slide, (910, 164, 1510, 650), "01-home.png")

    def build_slide_04():
        slide = base_native(4, "Practice Bank")
        title_native(slide, "Scale and definitions", "The numbers are counted by\nuse case.", "This slide prevents the common confusion between total drill pool, custom-test bank and revision modules.", width=760)
        metric_native(slide, (76, 395, 316, 555), "1,554*", "drill / daily pool", "583 vocabulary + 666 imbuhan + 305 persamaan")
        metric_native(slide, (346, 395, 586, 555), "971**", "custom-test bank", "666 imbuhan + 305 persamaan")
        metric_native(slide, (616, 395, 856, 555), "3", "revision surfaces", "Karangan phrases, surat rasmi and flashcards", COL["cyan"])
        add_text(slide, (76, 622, 860, 650), "* 583 + 666 + 305 = 1,554. ** 666 + 305 = 971.", 13.2, COL["ink"], True, "Aptos")
        add_text(slide, (76, 662, 846, 730), "Karangan, surat rasmi and flashcards are shown separately because they are revision surfaces, not custom-test question banks.", 10.2, COL["muted"])
        table_x, table_y = 940, 150
        rect(slide, (table_x, table_y, 1500, 680), COL["paper"], COL["line"], rounded=True, width=0.7)
        add_text(slide, (table_x + 28, table_y + 26, table_x + 220, table_y + 54), "Component", 10, COL["muted"], True)
        add_text(slide, (table_x + 248, table_y + 26, table_x + 350, table_y + 54), "Scale", 10, COL["muted"], True)
        add_text(slide, (table_x + 374, table_y + 26, table_x + 532, table_y + 54), "Use", 10, COL["muted"], True)
        rows = [
            ("Vocabulary", "583", "Drill / daily pool"),
            ("Imbuhan", "666", "Drill / daily + custom tests"),
            ("Persamaan", "305", "Drill / daily + custom tests"),
            ("Karangan", "297", "Writing phrases / revision"),
            ("Surat rasmi", "8 sections", "Templates and structure"),
            ("Flashcards", "23 sets", "Recall and review"),
        ]
        for i, row in enumerate(rows):
            y = table_y + 72 + i * 62
            rect(slide, (table_x + 24, y - 17, 1476, y - 16), COL["line"])
            add_text(slide, (table_x + 28, y, table_x + 225, y + 28), row[0], 10, COL["ink"], True)
            add_text(slide, (table_x + 248, y, table_x + 360, y + 28), row[1], 10, COL["ink"], True)
            add_text(slide, (table_x + 374, y, 1482, y + 28), row[2], 10, COL["muted"])

    def build_slide_05():
        slide = base_native(5, "Feature 1")
        title_native(slide, "Exam-skill modules", "Learners practise the exact\nweak area.", "This is not a generic language app: each section maps to an I2-relevant practice need.", width=660)
        browser_native(slide, (890, 144, 1510, 610), "04-persamaan.png")
        modules = [
            ("Vocabulary", "Meaning and usage\n583 items"),
            ("Imbuhan", "Word formation\n666 items"),
            ("Persamaan", "Synonym precision\n305 items"),
            ("Karangan", "Ready phrases\n297 phrases"),
            ("Surat rasmi", "Structure and tone\n8 sections"),
            ("Flashcards", "Recall and review\n23 sets"),
        ]
        for i, (h, b) in enumerate(modules):
            x = 76 + (i % 2) * 345
            y = 392 + (i // 2) * 120
            card_native(slide, (x, y, x + 305, y + 94), h, b, COL["lime"] if h in {"Imbuhan", "Persamaan"} else COL["cyan"])

    def build_slide_06():
        slide = base_native(6, "Feature 2")
        title_native(slide, "Focused custom drills", "A weak area becomes a\nshort test.", "Custom tests use the 971-item bank from imbuhan and persamaan question items.", width=700)
        browser_native(slide, (840, 132, 1510, 626), "06-test-setup.png")
        formula_native(slide, (76, 392, 740, 484), "Custom-test math", "666 imbuhan items + 305 persamaan items = 971 custom-test items.")
        steps = [
            ("Choose", "Select imbuhan or persamaan."),
            ("Set length", "Run 5, 10 or 20 questions depending on available time."),
            ("Review", "Use mistakes to decide what to repeat next."),
        ]
        for i, (h, b) in enumerate(steps):
            x = 76 + i * 255
            card_native(slide, (x, 540, x + 220, 690), h, b, COL["lime"] if i == 1 else COL["cyan"])

    def build_slide_07():
        slide = base_native(7, "Feature 3")
        title_native(slide, "Daily practice and usage", "Keep sessions small. Make\nbehaviour visible.", "Daily challenge draws from the 1,554-item pool; tracking records what users actually visit and practise.", width=760)
        browser_native(slide, (890, 132, 1510, 626), "07-daily-challenge.png")
        metric_native(slide, (76, 414, 286, 564), "10", "daily questions", "Small enough for a quick session")
        metric_native(slide, (316, 414, 526, 564), "1,554*", "daily pool", "Vocabulary + imbuhan + persamaan")
        metric_native(slide, (556, 414, 766, 564), "5-10", "minutes", "Designed around pockets of time")
        card_native(slide, (76, 622, 766, 746), "What can be observed", "Section visits, time spent by section, total session duration, active days and adoption across entities.", COL["cyan"])

    def build_slide_08():
        slide = base_native(8, "Demo and Value")
        title_native(slide, "How to present it", "Demo the learner flow, then\nshow the evidence.", "A clean 10-minute prep path, followed by the usage signal.", width=760)
        flow = [
            ("1", "Open dashboard", "Shows all practice paths in one place."),
            ("2", "Pick weak area", "Jump straight to vocabulary, imbuhan or persamaan."),
            ("3", "Run custom drill", "Turn a weak area into a short test."),
            ("4", "Check usage", "Show section and time tracking once enough data accumulates."),
        ]
        for i, (num, h, b) in enumerate(flow):
            y = 388 + i * 88
            rect(slide, (86, y, 126, y + 44), COL["lime"], rounded=True)
            add_text(slide, (86, y + 11, 126, y + 35), num, 10.5, COL["navy"], True, "Aptos", PP_ALIGN.CENTER)
            add_text(slide, (152, y, 700, y + 28), h, 13.5, COL["ink"], True, "Aptos")
            add_text(slide, (152, y + 30, 760, y + 58), b, 10.8, COL["muted"], False, "Aptos")
        browser_native(slide, (890, 132, 1510, 516), "01-home.png")
        metric_native(slide, (890, 570, 1060, 716), "40++", "learners", "Known adoption signal")
        metric_native(slide, (1085, 570, 1255, 716), "4+", "entities", "Cross-team use")
        metric_native(slide, (1280, 570, 1510, 716), "Ready", "analytics", "Usage evidence as logging matures")

    for builder in [
        build_slide_01,
        build_slide_02,
        build_slide_03,
        build_slide_04,
        build_slide_05,
        build_slide_06,
        build_slide_07,
        build_slide_08,
    ]:
        builder()

    prs.save(PPTX_PATH)


def main():
    if SLIDE_DIR.exists():
        shutil.rmtree(SLIDE_DIR)
    paths = [save_slide(factory(), name) for name, factory in SLIDES]
    build_ppt(paths)
    with ZipFile(PPTX_PATH) as z:
        slide_count = len([n for n in z.namelist() if n.startswith("ppt/slides/slide") and n.endswith(".xml")])
    print(PPTX_PATH)
    print(SLIDE_DIR)
    print(f"slides={slide_count}")


if __name__ == "__main__":
    main()
