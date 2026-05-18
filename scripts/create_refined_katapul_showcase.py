from pathlib import Path
from zipfile import ZipFile
import shutil

import numpy as np
from PIL import Image, ImageDraw, ImageFont, ImageFilter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "award-demo-assets"
PACKAGE = ROOT / "product-showcase-package"

DECK_DIR = PACKAGE / "katapult_modest_walkthrough_slides"
PPTX_PATH = PACKAGE / "KataPult_Modest_Product_Walkthrough_7slides.pptx"

FRAME_DIR = PACKAGE / "refined_teaser_frames"
GIF_PATH = PACKAGE / "KataPult_Modest_Product_Walkthrough.gif"
MP4_PATH = PACKAGE / "KataPult_Modest_Product_Walkthrough.mp4"
NOTES_PATH = PACKAGE / "KataPult_Modest_Product_Walkthrough_Notes.md"

W, H = 1600, 900
GW, GH = 1280, 720

COL = {
    "paper": "#F5F3EE",
    "white": "#FFFFFF",
    "ink": "#18202B",
    "muted": "#667085",
    "subtle": "#8A927D",
    "line": "#D8D3C8",
    "soft": "#ECE8DF",
    "soft2": "#F9F8F4",
    "green": "#7FAE2D",
    "green2": "#A7C957",
    "teal": "#1F6D78",
    "navy": "#111827",
    "amber": "#C88A2C",
}

COUNTS = {
    "Vocabulary": 583,
    "Imbuhan": 666,
    "Persamaan": 305,
    "Karangan": 297,
    "Surat Resmi": "8 sections + 10 examples",
    "Essay recall packs": "23 packs",
    "Daily pool": 1554,
    "Custom test bank": 971,
}


def font(size, weight="regular"):
    paths = {
        "display": ["C:/Windows/Fonts/bahnschrift.ttf", "C:/Windows/Fonts/segoeuil.ttf"],
        "regular": ["C:/Windows/Fonts/segoeui.ttf", "C:/Windows/Fonts/calibri.ttf"],
        "bold": ["C:/Windows/Fonts/segoeuib.ttf", "C:/Windows/Fonts/calibrib.ttf"],
    }
    key = "bold" if weight == "bold" else "display" if weight == "display" else "regular"
    for p in paths[key]:
        if Path(p).exists():
            return ImageFont.truetype(p, size)
    return ImageFont.load_default()


F = {
    "eyebrow": font(15, "bold"),
    "title": font(58, "display"),
    "h1": font(48, "display"),
    "h2": font(28, "bold"),
    "body": font(21, "regular"),
    "body_b": font(21, "bold"),
    "small": font(16, "regular"),
    "small_b": font(16, "bold"),
    "tiny": font(12, "regular"),
    "metric": font(42, "display"),
}

GF = {
    "eyebrow": font(13, "bold"),
    "hero": font(44, "display"),
    "h1": font(35, "display"),
    "h2": font(22, "bold"),
    "body": font(17, "regular"),
    "small": font(13, "regular"),
    "small_b": font(13, "bold"),
    "tiny": font(11, "regular"),
    "metric": font(28, "display"),
}


def rgb(hex_color):
    hex_color = hex_color.lstrip("#")
    return tuple(int(hex_color[i : i + 2], 16) for i in (0, 2, 4))


def rgba(hex_color, alpha=255):
    return rgb(hex_color) + (alpha,)


def clamp(v, lo=0, hi=1):
    return max(lo, min(hi, v))


def ease(t):
    t = clamp(t)
    return t * t * (3 - 2 * t)


def ease_out(t):
    t = clamp(t)
    return 1 - (1 - t) ** 3


def lerp(a, b, t):
    return a + (b - a) * t


def rounded(draw, box, fill, outline=None, radius=16, width=1):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def draw_text(draw, xy, value, fnt, fill, anchor=None):
    draw.text(xy, value, font=fnt, fill=fill, anchor=anchor)


def wrap_lines(draw, value, fnt, max_width):
    lines, cur = [], ""
    for word in str(value).split():
        candidate = (cur + " " + word).strip()
        if draw.textbbox((0, 0), candidate, font=fnt)[2] <= max_width:
            cur = candidate
        else:
            if cur:
                lines.append(cur)
            cur = word
    if cur:
        lines.append(cur)
    return lines


def para(draw, xy, value, fnt, fill, max_width, gap=6):
    x, y = xy
    for line in wrap_lines(draw, value, fnt, max_width):
        draw_text(draw, (x, y), line, fnt, fill)
        y += fnt.size + gap
    return y


def shadow(layer, box, blur=14, alpha=30, radius=18):
    sh = Image.new("RGBA", layer.size, (0, 0, 0, 0))
    sd = ImageDraw.Draw(sh, "RGBA")
    sd.rounded_rectangle(box, radius=radius, fill=(30, 30, 30, alpha))
    sh = sh.filter(ImageFilter.GaussianBlur(blur))
    layer.alpha_composite(sh)


def load_shot(name, size, fill=COL["white"]):
    im = Image.open(ASSETS / name).convert("RGBA")
    if im.height > 780:
        im = im.crop((0, 0, im.width, 780))
    im.thumbnail(size)
    canvas = Image.new("RGBA", size, rgba(fill))
    canvas.alpha_composite(im, ((size[0] - im.width) // 2, (size[1] - im.height) // 2))
    return canvas


def browser_preview(layer, box, screenshot, radius=20):
    shadow(layer, box, blur=16, alpha=34, radius=radius)
    d = ImageDraw.Draw(layer, "RGBA")
    rounded(d, box, COL["white"], outline=COL["line"], radius=radius, width=1)
    x1, y1, x2, y2 = box
    rounded(d, (x1, y1, x2, y1 + 40), "#E8ECEF", radius=radius)
    for i, c in enumerate(["#E36F66", "#E2B646", "#55B983"]):
        d.ellipse((x1 + 22 + i * 24, y1 + 14, x1 + 34 + i * 24, y1 + 26), fill=c)
    d.rounded_rectangle((x1 + 146, y1 + 14, x2 - 28, y1 + 27), radius=6, fill="#D7DEE6")
    shot = load_shot(screenshot, (x2 - x1 - 26, y2 - y1 - 52))
    layer.alpha_composite(shot, (x1 + 13, y1 + 46))


def base_slide(slide_no, section):
    img = Image.new("RGBA", (W, H), rgba(COL["paper"]))
    d = ImageDraw.Draw(img, "RGBA")
    d.line((72, 74, W - 72, 74), fill=COL["line"], width=1)
    draw_text(d, (72, 42), "KataPult", F["small_b"], COL["ink"])
    draw_text(d, (172, 42), section, F["small"], COL["subtle"])
    draw_text(d, (W - 72, 42), f"{slide_no:02d}", F["small_b"], COL["subtle"], anchor="ra")
    return img


def title_block(draw, eyebrow, title, subtitle=None, x=72, y=126, width=700):
    draw_text(draw, (x, y), eyebrow.upper(), F["eyebrow"], COL["green"])
    title_end = para(draw, (x, y + 36), title, F["title"], COL["ink"], width, gap=4)
    if subtitle:
        para(draw, (x, title_end + 18), subtitle, F["body"], COL["muted"], width, gap=6)


def small_card(draw, box, title, body=None, accent=None):
    rounded(draw, box, COL["white"], outline=COL["line"], radius=14, width=1)
    if accent:
        draw.line((box[0] + 18, box[1] + 18, box[0] + 18, box[3] - 18), fill=accent, width=5)
        x = box[0] + 40
    else:
        x = box[0] + 22
    draw_text(draw, (x, box[1] + 20), title, F["small_b"], COL["ink"])
    if body:
        para(draw, (x, box[1] + 48), body, F["small"], COL["muted"], box[2] - x - 20, gap=4)


def metric(draw, box, value, label, note=None, accent=COL["green"]):
    rounded(draw, box, COL["white"], outline=COL["line"], radius=14, width=1)
    draw_text(draw, (box[0] + 20, box[1] + 16), value, F["metric"], accent)
    draw_text(draw, (box[0] + 20, box[1] + 62), label, F["small_b"], COL["ink"])
    if note:
        para(draw, (box[0] + 20, box[1] + 90), note, F["tiny"], COL["muted"], box[2] - box[0] - 40, gap=3)


def card_label(draw, box, title, body, accent=COL["green"]):
    rounded(draw, box, COL["white"], outline=COL["line"], radius=14, width=1)
    draw.line((box[0] + 16, box[1] + 18, box[0] + 16, box[3] - 18), fill=accent, width=5)
    draw_text(draw, (box[0] + 36, box[1] + 18), title, F["body_b"], COL["ink"])
    if body:
        para(draw, (box[0] + 36, box[1] + 52), body, F["small"], COL["muted"], box[2] - box[0] - 60, gap=4)


def arrow(draw, start, end, color=COL["green"]):
    x1, y1 = start
    x2, y2 = end
    draw.line((x1, y1, x2, y2), fill=color, width=5)
    draw.polygon([(x2, y2), (x2 - 16, y2 - 9), (x2 - 16, y2 + 9)], fill=color)


def slide_cover():
    img = base_slide(1, "Practice Companion")
    d = ImageDraw.Draw(img, "RGBA")
    title_block(
        d,
        "Working prototype",
        "KataPult: Bahasa Indo\nPractice Companion",
        "A working prototype for short, structured Bahasa Indo revision.",
        width=720,
    )
    browser_preview(img, (790, 128, 1510, 694), "01-home.png")
    rounded(d, (76, 740, 1510, 814), COL["white"], outline=COL["line"], radius=14)
    draw_text(
        d,
        (116, 766),
        "Relevant Bahasa Indo learning materials centralised into one practice companion.",
        F["body_b"],
        COL["ink"],
    )
    return img


def slide_problem():
    img = base_slide(2, "Problem")
    d = ImageDraw.Draw(img, "RGBA")
    title_block(
        d,
        "Preparation reality",
        "The hard part was not\nfinding content. It was\nrestarting practice.",
        "Learners had notes, screenshots, online references, recall packs and generated practice material, but no clear place to continue after a break.",
        width=830,
    )
    rounded(d, (930, 174, 1468, 372), COL["white"], outline=COL["line"], radius=16)
    draw_text(d, (958, 200), "Scattered inputs", F["body_b"], COL["ink"])
    inputs = ["Notes", "Screenshots", "Online references", "Recall packs", "Generated practice material"]
    for i, lab in enumerate(inputs):
        y = 242 + i * 24
        d.ellipse((960, y + 5, 968, y + 13), fill=COL["green"])
        draw_text(d, (982, y), lab, F["small"], COL["muted"])
    d.line((1199, 382, 1199, 430), fill=COL["green"], width=4)
    d.polygon([(1199, 446), (1187, 426), (1211, 426)], fill=COL["green"])
    rounded(d, (930, 456, 1468, 532), COL["soft2"], outline=COL["line"], radius=16)
    draw_text(d, (1199, 494), "\"Where do I continue?\"", F["body_b"], COL["ink"], anchor="mm")
    d.line((1199, 542, 1199, 590), fill=COL["green"], width=4)
    d.polygon([(1199, 606), (1187, 586), (1211, 586)], fill=COL["green"])
    rounded(d, (930, 618, 1468, 694), COL["soft"], outline=COL["line"], radius=16)
    draw_text(d, (1199, 656), "Too much setup before practice", F["body_b"], COL["ink"], anchor="mm")
    pains = [
        ("No clear next step", "Each session started with deciding what to revise."),
        ("Weak areas were hard to isolate", "Imbuhan, persamaan and formal writing need repeated practice."),
        ("Progress was hard to see", "Attempts, mistakes and return points were not easy to track."),
    ]
    for i, (h, b) in enumerate(pains):
        card_label(d, (76, 458 + i * 104, 760, 540 + i * 104), h, b, COL["green"])
    return img


def slide_learning_jobs():
    img = base_slide(3, "What Changes")
    d = ImageDraw.Draw(img, "RGBA")
    title_block(
        d,
        "What KataPult changes",
        "Scattered materials\nbecome short\npractice flows.",
        "KataPult centralises relevant learning materials in one place, then routes them into drills, tests, feedback and writing-support pages.",
        width=760,
    )
    left = ["Notes", "Online references", "Screenshots", "Recall packs", "Generated practice material"]
    right = ["Daily challenge", "Component drills", "Custom tests", "Feedback explanations", "Writing guides and examples"]
    for i, lab in enumerate(left):
        y = 416 + i * 50
        d.rounded_rectangle((78, y, 420, y + 38), radius=10, fill=COL["white"], outline=COL["line"])
        draw_text(d, (98, y + 10), lab, F["small_b"], COL["ink"])
    rounded(d, (562, 474, 964, 606), COL["white"], outline=COL["line"], radius=16)
    draw_text(d, (763, 512), "KataPult", F["h2"], COL["green"], anchor="mm")
    draw_text(d, (763, 552), "organises content", F["body_b"], COL["ink"], anchor="mm")
    arrow(d, (440, 540), (538, 540), COL["green"])
    arrow(d, (988, 540), (1086, 540), COL["green"])
    for i, lab in enumerate(right):
        y = 416 + i * 50
        d.rounded_rectangle((1112, y, 1490, y + 38), radius=10, fill="#EEF6E8", outline=COL["line"])
        draw_text(d, (1132, y + 10), lab, F["small_b"], COL["ink"])
    rounded(d, (78, 752, 1490, 810), COL["soft2"], outline=COL["line"], radius=14)
    draw_text(
        d,
        (106, 772),
        "Review note: AI-assisted practice items are flagged for human review before being treated as validated material.",
        F["small_b"],
        COL["muted"],
    )
    return img


def slide_product_loop():
    img = base_slide(4, "Product Loop")
    d = ImageDraw.Draw(img, "RGBA")
    title_block(
        d,
        "Practice loop",
        "A simple loop keeps\npractice moving.",
        "Learners can pick a module, practise briefly, understand the answer and return through daily challenge, XP or streaks.",
        width=700,
    )
    loop = [
        ("1", "Pick a module", "Start with the area that needs work."),
        ("2", "Practise briefly", "Use a daily challenge or focused test."),
        ("3", "Get feedback", "Right/wrong + LLM explanation."),
        ("4", "Review mistakes", "Repeat missed items or continue."),
        ("5", "Return later", "Use daily challenge, XP or streaks."),
    ]
    y = 604
    for i, (n, h, b) in enumerate(loop):
        x = 76 + i * 296
        rounded(d, (x, y, x + 246, y + 148), COL["white"], outline=COL["line"], radius=16)
        rounded(d, (x + 22, y + 22, x + 58, y + 58), COL["green"], radius=10)
        draw_text(d, (x + 40, y + 41), n, F["small_b"], COL["white"], anchor="mm")
        draw_text(d, (x + 74, y + 24), h, F["small_b"], COL["ink"])
        para(d, (x + 74, y + 56), b, F["tiny"], COL["muted"], 146, gap=3)
        if i < 4:
            arrow(d, (x + 252, y + 74), (x + 286, y + 74), COL["line"])
    browser_preview(img, (808, 140, 1510, 520), "07-daily-challenge.png")
    small_card(d, (808, 536, 1010, 588), "Start a short drill", None, COL["green"])
    small_card(d, (1034, 536, 1240, 588), "Track progress", None, COL["teal"])
    small_card(d, (1264, 536, 1510, 588), "Understand answers", None, COL["green"])
    return img


def slide_feature_walkthrough():
    img = base_slide(5, "Feature Walkthrough")
    d = ImageDraw.Draw(img, "RGBA")
    title_block(
        d,
        "Feature walkthrough",
        "Each component has\na specific learning job.",
        "The app separates browsing, typed practice, recall, writing guides and focused tests so learners do not revise from one giant list.",
        width=760,
    )
    rows = [
        ("Vocabulary", "Meanings, examples and translations for core word building."),
        ("Imbuhan", "Affix practice: apply prefixes, suffixes and word-form rules correctly."),
        ("Persamaan MCQ", "Synonym recognition through sentence-context questions."),
        ("Persamaan Latihan", "Typed synonym practice with immediate feedback."),
        ("Karangan", "Essay vocabulary, phrase banks and definitions for composition writing."),
        ("Essay Recall Packs", "Fast recall of openings, conclusions, examples and useful essay phrases."),
        ("Surat Resmi", "Formal letter structure, tone, standard phrases and sample responses."),
        ("Custom Test", "Configurable drills for focused revision on selected components."),
        ("Progress", "Streaks, XP and levels that make return practice visible."),
    ]
    x1, x2 = 76, 924
    y0, row_h = 384, 44
    rounded(d, (x1, y0 - 18, x2, y0 + row_h * len(rows) + 18), COL["white"], outline=COL["line"], radius=16)
    for i, (h, b) in enumerate(rows):
        y = y0 + i * row_h
        if i:
            d.line((x1 + 24, y - 8, x2 - 24, y - 8), fill=COL["line"], width=1)
        draw_text(d, (x1 + 28, y), h, F["small_b"], COL["ink"])
        draw_text(d, (x1 + 286, y), b, F["small"], COL["muted"])
    browser_preview(img, (1000, 272, 1510, 720), "03-imbuhan.png")
    small_card(d, (1000, 724, 1510, 822), "Shown action", "Learner answers, checks feedback, and can read an LLM-generated explanation.", COL["green"])
    return img


def slide_current_scale():
    img = base_slide(6, "Current State")
    d = ImageDraw.Draw(img, "RGBA")
    title_block(
        d,
        "Current scale",
        "Current state: useful enough to trial,\nstill early enough to improve.",
        "The current numbers show coverage and informal trial use. Stronger learning claims should come after content review and usage analysis.",
        width=900,
    )
    metrics = [
        ("1,554", "active drill-ready items", "Vocabulary, imbuhan and persamaan question pool.", COL["green"]),
        ("971", "current custom-test items", "Custom Test currently covers imbuhan + persamaan.", COL["teal"]),
        ("10++", "early informal trial users", "Enough early use to keep improving.", COL["amber"]),
        ("Multiple", "practice modes", "Drills, recall, writing guides and tests.", COL["green"]),
    ]
    for i, (v, lab, note, col) in enumerate(metrics):
        x = 76 + (i % 2) * 442
        y = 414 + (i // 2) * 164
        metric(d, (x, y, x + 390, y + 126), v, lab, note, col)
    browser_preview(img, (1000, 226, 1510, 620), "06-test-setup.png")
    rounded(d, (76, 748, 1510, 812), COL["soft2"], outline=COL["line"], radius=14)
    draw_text(
        d,
        (106, 768),
        "Footnote: These numbers describe content coverage and informal trial usage. Learning impact should be validated through future review.",
        F["small_b"],
        COL["muted"],
    )
    return img


def slide_next_steps():
    img = base_slide(7, "Next Steps")
    d = ImageDraw.Draw(img, "RGBA")
    title_block(
        d,
        "Next refinement steps",
        "The app helps learners\nrestart practice more\neasily.",
        "It centralises relevant learning materials, gives learners a clear place to continue, and explains answers after practice.",
        width=850,
    )
    steps = [
        ("Content quality", "Review collected items and generated explanations before treating them as validated material.", COL["green"]),
        ("User feedback", "Collect what learners find useful, confusing or missing during real use.", COL["teal"]),
        ("Iterative refinement", "Use feedback and usage patterns to improve the most-used modules first.", COL["amber"]),
    ]
    for i, (h, b, col) in enumerate(steps):
        x = 76 + i * 486
        card_label(d, (x, 430, x + 424, 592), h, b, col)
    rounded(d, (76, 690, 1510, 796), COL["white"], outline=COL["line"], radius=16)
    para(
        d,
        (112, 720),
        "KataPult should keep improving through continuous user feedback, content review and small product refinements based on real practice behaviour.",
        F["h2"],
        COL["ink"],
        1340,
        gap=4,
    )
    return img


def slide_demo_measurement():
    img = base_slide(7, "Demo and Measurement")
    d = ImageDraw.Draw(img, "RGBA")
    title_block(
        d,
        "Demo plan",
        "A short demo,\nwith measured follow-up.",
        "Keep it practical: one learner, one weak area, one short practice session.",
        width=760,
    )
    steps = [
        ("1", "Open dashboard", "Show modules, streak and recommended path."),
        ("2", "Pick a weak area", "Use imbuhan or persamaan to make the flow concrete."),
        ("3", "Run a short drill", "Use 5 or 10 questions; show feedback."),
        ("4", "Review flashcards or phrases", "Show how revision continues outside test mode."),
        ("5", "Check tracking", "Show section visits and time spent once data accumulates."),
    ]
    for i, (n, h, b) in enumerate(steps):
        y = 348 + i * 70
        rounded(d, (86, y, 124, y + 38), COL["green"], radius=10)
        draw_text(d, (105, y + 19), n, F["small_b"], COL["white"], anchor="mm")
        draw_text(d, (150, y - 2), h, F["body_b"], COL["ink"])
        draw_text(d, (150, y + 29), b, F["small"], COL["muted"])
    browser_preview(img, (890, 146, 1510, 542), "01-home.png")
    small_card(d, (890, 596, 1510, 720), "Use evidence carefully", "Report what users actually do after a defined review window, then decide what to improve next.", COL["green"])
    return img


DECK_SLIDES = [
    ("01_cover", slide_cover),
    ("02_problem", slide_problem),
    ("03_learning_jobs", slide_learning_jobs),
    ("04_product_loop", slide_product_loop),
    ("05_feature_walkthrough", slide_feature_walkthrough),
    ("06_current_scale", slide_current_scale),
    ("07_next_steps", slide_next_steps),
]


def save_deck_previews():
    if DECK_DIR.exists():
        shutil.rmtree(DECK_DIR)
    DECK_DIR.mkdir(parents=True, exist_ok=True)
    paths = []
    for name, factory in DECK_SLIDES:
        path = DECK_DIR / f"{name}.png"
        factory().convert("RGB").save(path, quality=95)
        paths.append(path)
    return paths


class PptBuilder:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width = Inches(13.333333)
        self.prs.slide_height = Inches(7.5)
        self.blank = self.prs.slide_layouts[6]

    def sx(self, px):
        return Inches(px / W * 13.333333)

    def sy(self, px):
        return Inches(px / H * 7.5)

    def pbox(self, box):
        x1, y1, x2, y2 = box
        return self.sx(x1), self.sy(y1), self.sx(x2 - x1), self.sy(y2 - y1)

    def color(self, value):
        return RGBColor(*rgb(value))

    def slide(self, number, section):
        s = self.prs.slides.add_slide(self.blank)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = self.color(COL["paper"])
        self.line(s, (72, 74, W - 72, 74), COL["line"])
        self.text(s, (72, 42, 170, 64), "KataPult", 10.5, COL["ink"], True)
        self.text(s, (172, 42, 660, 64), section, 10, COL["subtle"])
        self.text(s, (1470, 42, W - 72, 64), f"{number:02d}", 10, COL["subtle"], True, align=PP_ALIGN.RIGHT)
        return s

    def shape(self, slide, box, fill, line=None, rounded=False, width=0.7):
        shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, *self.pbox(box))
        shp.fill.solid()
        shp.fill.fore_color.rgb = self.color(fill)
        if line:
            shp.line.color.rgb = self.color(line)
            shp.line.width = Pt(width)
        else:
            shp.line.fill.background()
        return shp

    def line(self, slide, box, color, width=1):
        x1, y1, x2, y2 = box
        shp = slide.shapes.add_connector(1, self.sx(x1), self.sy(y1), self.sx(x2), self.sy(y2))
        shp.line.color.rgb = self.color(color)
        shp.line.width = Pt(width)
        return shp

    def oval(self, slide, box, fill):
        shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, *self.pbox(box))
        shp.fill.solid()
        shp.fill.fore_color.rgb = self.color(fill)
        shp.line.fill.background()
        return shp

    def text(self, slide, box, value, size, color=COL["ink"], bold=False, font_name="Segoe UI", align=None):
        tx = slide.shapes.add_textbox(*self.pbox(box))
        tf = tx.text_frame
        tf.clear()
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = 0
        tf.margin_right = 0
        tf.margin_top = 0
        tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = str(value)
        p.space_after = Pt(0)
        p.line_spacing = 1.05
        if align is not None:
            p.alignment = align
        for run in p.runs:
            run.font.name = font_name
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = self.color(color)
        return tx

    def title(self, slide, eyebrow, title, subtitle=None, x=72, y=126, width=700):
        self.text(slide, (x, y, x + width, y + 24), eyebrow.upper(), 9, COL["green"], True, "Segoe UI")
        lines = title.count("\n") + 1
        self.text(slide, (x, y + 36, x + width, y + 36 + lines * 62), title, 34, COL["ink"], False, "Bahnschrift")
        if subtitle:
            self.text(slide, (x, y + 54 + lines * 62, x + width, y + 138 + lines * 62), subtitle, 13.5, COL["muted"])

    def add_picture_fit(self, slide, path, box):
        src = Path(path)
        im = Image.open(src)
        if im.height > 780:
            crop_dir = PACKAGE / "refined_screenshot_crops"
            crop_dir.mkdir(parents=True, exist_ok=True)
            cropped = crop_dir / src.name
            im = im.crop((0, 0, im.width, 780))
            im.save(cropped)
            src = cropped
        max_w, max_h = box[2] - box[0], box[3] - box[1]
        ratio = min(max_w / im.width, max_h / im.height)
        pw, ph = im.width * ratio, im.height * ratio
        left = box[0] + (max_w - pw) / 2
        top = box[1] + (max_h - ph) / 2
        slide.shapes.add_picture(str(src), self.sx(left), self.sy(top), width=self.sx(pw), height=self.sy(ph))

    def browser(self, slide, box, screenshot):
        self.shape(slide, (box[0] + 8, box[1] + 12, box[2] + 8, box[3] + 12), "#E1DDD2", rounded=True)
        self.shape(slide, box, COL["white"], COL["line"], rounded=True)
        self.shape(slide, (box[0], box[1], box[2], box[1] + 40), "#E8ECEF", rounded=True)
        for i, c in enumerate(["#E36F66", "#E2B646", "#55B983"]):
            self.oval(slide, (box[0] + 22 + i * 24, box[1] + 14, box[0] + 34 + i * 24, box[1] + 26), c)
        self.shape(slide, (box[0] + 146, box[1] + 14, box[2] - 28, box[1] + 27), "#D7DEE6", rounded=True)
        self.add_picture_fit(slide, ASSETS / screenshot, (box[0] + 13, box[1] + 46, box[2] - 13, box[3] - 6))

    def card(self, slide, box, title, body=None, accent=None):
        self.shape(slide, box, COL["white"], COL["line"], rounded=True)
        x = box[0] + 22
        if accent:
            self.shape(slide, (box[0] + 18, box[1] + 18, box[0] + 23, box[3] - 18), accent, rounded=True)
            x = box[0] + 40
        self.text(slide, (x, box[1] + 20, box[2] - 20, box[1] + 48), title, 10.5, COL["ink"], True)
        if body:
            self.text(slide, (x, box[1] + 50, box[2] - 18, box[3] - 16), body, 9.2, COL["muted"])

    def metric(self, slide, box, value, label, note=None, accent=COL["green"]):
        self.shape(slide, box, COL["white"], COL["line"], rounded=True)
        self.text(slide, (box[0] + 20, box[1] + 16, box[2] - 20, box[1] + 60), value, 24, accent, False, "Bahnschrift")
        self.text(slide, (box[0] + 20, box[1] + 64, box[2] - 20, box[1] + 90), label, 10, COL["ink"], True)
        if note:
            self.text(slide, (box[0] + 20, box[1] + 94, box[2] - 20, box[3] - 14), note, 7.5, COL["muted"])

    def save(self, path):
        self.prs.save(path)


def build_ppt():
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    for name, _ in DECK_SLIDES:
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(DECK_DIR / f"{name}.png"), 0, 0, width=prs.slide_width, height=prs.slide_height)
    prs.save(PPTX_PATH)


def save_notes_blueprint():
    NOTES_PATH.write_text(
        """# KataPult Modest Product Walkthrough Notes

## Slide 1: Title
Title: KataPult: Bahasa Indo Practice Companion
Subtitle: A working prototype for short, structured Bahasa Indo revision.
Screenshot guidance: Use the homepage screenshot large enough to prove this is a real app, without asking viewers to read every UI detail.
Speaker notes: Open with a modest frame. KataPult is a working prototype that centralises relevant Bahasa Indo learning materials into one practice companion.

## Slide 2: Problem
Headline: The hard part was not finding content. It was restarting practice.
Screenshot/visual guidance: Use the clean vertical before-state flow: scattered inputs, "Where do I continue?", then too much setup before practice.
Speaker notes: The main pain is decision load. Learners may have useful material, but they still need to decide where to continue, what to revise and how to test themselves.

## Slide 3: What KataPult Changes
Headline: Scattered materials become short practice flows.
Screenshot/visual guidance: Use an input-to-output flow, with the review note kept as a slim bottom note.
Speaker notes: KataPult does not replace learning material. It centralises relevant notes, references, recall packs and generated practice material, then routes them into daily challenge, component drills, custom tests, feedback explanations and writing guides and examples.

## Slide 4: Practice Loop
Headline: A simple loop keeps practice moving.
Screenshot guidance: Use Daily Challenge as the main screenshot, annotated with simple user actions.
Speaker notes: The flow gives learners a practical way to continue: pick a module, practise briefly, get right/wrong feedback, read an LLM-generated explanation where available, review mistakes or continue, then return later through daily challenge, XP or streaks.

## Slide 5: Feature Walkthrough
Headline: Each component has a specific learning job.
Screenshot guidance: Use Imbuhan as the module screenshot because it shows typed practice, hints, checking and review clearly.
Speaker notes: Vocabulary covers meanings, examples and translations. Imbuhan supports affix and word-form practice. Persamaan MCQ supports synonym recognition in context. Persamaan Latihan supports typed synonym practice with immediate feedback. Karangan supports essay vocabulary, phrase banks and definitions. Essay Recall Packs help fast recall of openings, conclusions, examples and useful phrases. Surat Resmi covers structure, tone, standard phrases and sample responses. Custom Test provides configurable focused drills. Progress indicators make return practice visible.

## Slide 6: Current Scale And Early Usage Signal
Headline: Current state: useful enough to trial, still early enough to improve.
Screenshot guidance: Use Custom Test setup so the 971-item test scope is grounded in a product view.
Speaker notes: Use the numbers carefully. 1,554 active drill-ready items and 971 current custom-test items describe coverage. 10++ early informal trial users show enough early use to keep improving the product. Stronger learning claims should come after content review and usage analysis.

## Slide 7: Next Refinement Steps
Headline: The app helps learners restart practice more easily.
Screenshot guidance: No screenshot required; keep the close simple and credible.
Speaker notes: Close with the current product value: the app centralises relevant materials, gives learners a clearer place to continue and helps them understand answers through feedback and explanations. The next step is continuous improvement through user feedback, content review and small product refinements based on real practice behaviour.

## Tone Cleanup Table
| Old phrase | Replacement |
| --- | --- |
| flashcard sets | essay recall packs / topic and phrase packs |
| writing scaffolds | writing guides and examples |
| prompt outputs | generated practice material |
| LLM-generated expansions | AI-assisted practice items, flagged for review |
| 4+ teams | 10++ early users have trialed the app informally |
| impact | early usage signal |
| behaviour change | easier restart point / repeatable practice loop |
| master Bahasa Indonesia | practise Bahasa Indo |
| adopted across teams | trialed informally |

## QA Checklist Used
- Slide 1 makes the app understandable in under 10 seconds.
- Slide 2 explains the restart problem without product jargon.
- Slide 3 shows inputs becoming practice outputs without cropped boxes.
- Slide 4 explains the repeatable practice loop and answer explanations in product language.
- Slide 5 shows breadth without a dense grid of similar cards.
- Slide 6 treats metrics as coverage and early trial usage.
- Slide 7 closes with practical next steps.
- Screenshots are larger and annotated only where they demonstrate a user action.
- Cards align to a strict grid and avoid text clipping.
- Tone stays modest, clear and senior-leader safe.
""",
        encoding="utf-8",
    )
    return NOTES_PATH


def create_contact_sheet(paths):
    thumb_w, thumb_h = 400, 225
    pad = 24
    cols = 2
    rows = (len(paths) + cols - 1) // cols
    out = Image.new("RGB", (cols * thumb_w + (cols + 1) * pad, rows * (thumb_h + 34) + (rows + 1) * pad), "white")
    d = ImageDraw.Draw(out)
    for i, p in enumerate(paths):
        im = Image.open(p).convert("RGB")
        im.thumbnail((thumb_w, thumb_h))
        x = pad + (i % cols) * (thumb_w + pad)
        y = pad + (i // cols) * (thumb_h + 34 + pad)
        out.paste(im, (x, y))
        d.text((x, y + thumb_h + 8), p.name, fill=(72, 72, 72))
    out_path = DECK_DIR / "_contact_sheet.png"
    out.save(out_path, quality=95)
    return out_path


def gif_bg():
    img = Image.new("RGBA", (GW, GH), rgba(COL["paper"]))
    d = ImageDraw.Draw(img, "RGBA")
    for x in range(0, GW, 160):
        d.line((x, 0, x, GH), fill=rgba(COL["line"], 65), width=1)
    d.rectangle((0, 0, GW, 68), fill=rgba(COL["paper"]))
    d.line((58, 66, GW - 58, 66), fill=COL["line"], width=1)
    draw_text(d, (58, 34), "KataPult", GF["small_b"], COL["ink"])
    draw_text(d, (150, 34), "Bahasa Indo Prep Companion", GF["small"], COL["subtle"])
    return img


def g_text(draw, xy, value, fnt, fill, max_width=None):
    if max_width:
        return para(draw, xy, value, fnt, fill, max_width, gap=5)
    draw_text(draw, xy, value, fnt, fill)
    return xy[1] + fnt.size


def g_browser(layer, box, screenshot, scale=1.0):
    x1, y1, x2, y2 = box
    cx, cy = (x1 + x2) / 2, (y1 + y2) / 2
    bw, bh = (x2 - x1) * scale, (y2 - y1) * scale
    x1, y1, x2, y2 = int(cx - bw / 2), int(cy - bh / 2), int(cx + bw / 2), int(cy + bh / 2)
    shadow(layer, (x1, y1, x2, y2), blur=14, alpha=38, radius=18)
    d = ImageDraw.Draw(layer, "RGBA")
    rounded(d, (x1, y1, x2, y2), COL["white"], outline=COL["line"], radius=18)
    rounded(d, (x1, y1, x2, y1 + 34), "#E8ECEF", radius=18)
    for i, c in enumerate(["#E36F66", "#E2B646", "#55B983"]):
        d.ellipse((x1 + 18 + i * 20, y1 + 12, x1 + 29 + i * 20, y1 + 23), fill=c)
    d.rounded_rectangle((x1 + 124, y1 + 12, x2 - 24, y1 + 23), radius=5, fill="#D7DEE6")
    shot = load_shot(screenshot, (x2 - x1 - 22, y2 - y1 - 46))
    layer.alpha_composite(shot, (x1 + 11, y1 + 40))


def g_note(draw, box, title, body=None, accent=COL["green"]):
    rounded(draw, box, COL["white"], outline=COL["line"], radius=13)
    draw.line((box[0] + 16, box[1] + 16, box[0] + 16, box[3] - 16), fill=accent, width=4)
    draw_text(draw, (box[0] + 34, box[1] + 14), title, GF["small_b"], COL["ink"])
    if body:
        para(draw, (box[0] + 34, box[1] + 38), body, GF["small"], COL["muted"], box[2] - box[0] - 50, gap=3)


def scene_problem(p):
    img = gif_bg()
    d = ImageDraw.Draw(img, "RGBA")
    x_shift = int((1 - ease_out(p)) * -24)
    g_text(d, (58 + x_shift, 112), "Material existed.", GF["hero"], COL["ink"])
    g_text(d, (58 + x_shift, 168), "The issue was restart friction.", GF["body"], COL["muted"], 480)
    labels = ["Online refs", "Notes", "Screenshots", "Recall packs", "Generated material", "Raw lists"]
    for i, lab in enumerate(labels):
        reveal = ease_out(p * 2.4 - i * 0.18)
        y = 292 + i * 44
        x = int(62 + (1 - reveal) * -32)
        alpha = int(255 * clamp(reveal))
        d.rounded_rectangle((x, y, x + 210, y + 30), radius=9, fill=rgba(COL["white"], alpha), outline=rgba(COL["line"], alpha))
        d.text((x + 14, y + 7), lab, font=GF["small_b"], fill=rgba(COL["ink"], alpha))
    g_note(d, (524, 298, 1110, 420), "Before", "Useful material was scattered across references, notes, recall packs and generated practice material.", COL["teal"])
    g_note(d, (524, 450, 1110, 572), "KataPult response", "Organise it around daily restart, targeted drills, writing support and assessment.", COL["green"])
    return img


def scene_modules(p):
    img = gif_bg()
    d = ImageDraw.Draw(img, "RGBA")
    g_text(d, (58, 112), "Practice by learning job.", GF["hero"], COL["ink"], 520)
    g_text(d, (58, 222), "Daily restart, targeted drills, writing support and assessment.", GF["body"], COL["muted"], 450)
    scale = 0.96 + 0.025 * ease(p)
    g_browser(img, (542, 108, 1190, 596), "01-home.png", scale=scale)
    items = [("Vocabulary", "583"), ("Imbuhan", "666"), ("Persamaan", "305"), ("Karangan", "297")]
    for i, (label, count) in enumerate(items):
        y = 346 + i * 56
        g_note(d, (58, y, 390, y + 42), f"{label}   {count}", None, COL["green"] if i < 3 else COL["teal"])
    return img


def scene_drill(p):
    img = gif_bg()
    d = ImageDraw.Draw(img, "RGBA")
    g_text(d, (58, 112), "Choose, practise,\nreview, return.", GF["hero"], COL["ink"], 500)
    g_text(d, (58, 238), "A clear place to continue and build learning momentum.", GF["body"], COL["muted"], 440)
    g_note(d, (58, 320, 442, 410), "Custom-test bank", "666 imbuhan + 305 persamaan = 971 items.", COL["green"])
    g_note(d, (58, 438, 442, 528), "Session shape", "Choose component, set length, answer, review.", COL["teal"])
    g_browser(img, (548, 120, 1188, 600), "06-test-setup.png", scale=0.96 + 0.02 * ease(p))
    return img


def scene_daily(p):
    img = gif_bg()
    d = ImageDraw.Draw(img, "RGBA")
    g_text(d, (58, 112), "Current scale supports\nfurther refinement.", GF["hero"], COL["ink"], 540)
    g_text(d, (58, 244), "Coverage and early use give clear next refinement steps.", GF["body"], COL["muted"], 450)
    g_note(d, (58, 334, 240, 414), "10", "daily questions", COL["green"])
    g_note(d, (260, 334, 470, 414), "1,554*", "drill / daily pool", COL["green"])
    g_note(d, (58, 446, 470, 538), "* 583 + 666 + 305", "Vocabulary, imbuhan and persamaan feed daily practice.", COL["teal"])
    g_browser(img, (552, 112, 1190, 604), "07-daily-challenge.png", scale=0.96 + 0.018 * ease(p))
    return img


def scene_close(p):
    img = gif_bg()
    d = ImageDraw.Draw(img, "RGBA")
    g_text(d, (58, 124), "KataPult", GF["hero"], COL["ink"])
    g_text(d, (58, 184), "Useful now. Ready for quality and usage review.", GF["h2"], COL["teal"])
    steps = ["Validate content quality", "Learn from actual usage", "Refine the product"]
    for i, step in enumerate(steps):
        y = 300 + i * 62
        d.rounded_rectangle((64, y, 94, y + 30), radius=8, fill=COL["green"])
        d.text((79, y + 15), str(i + 1), font=GF["small_b"], fill=COL["white"], anchor="mm")
        d.text((114, y + 6), step, font=GF["h2"], fill=COL["ink"])
    g_browser(img, (590, 116, 1196, 574), "01-home.png", scale=0.94)
    d.text((58, 638), "Next: review AI-assisted items and collect clearer usage evidence.", font=GF["tiny"], fill=COL["muted"])
    return img


GIF_SCENES = [
    (0.0, 2.0, scene_problem),
    (2.0, 4.1, scene_modules),
    (4.1, 6.0, scene_drill),
    (6.0, 7.8, scene_daily),
    (7.8, 9.2, scene_close),
]


def render_gif_at(t):
    for i, (start, end, painter) in enumerate(GIF_SCENES):
        if start <= t <= end or i == len(GIF_SCENES) - 1:
            p = clamp((t - start) / (end - start))
            frame = painter(p)
            trans = 0.18
            if i < len(GIF_SCENES) - 1 and end - t < trans:
                nxt = GIF_SCENES[i + 1][2](0)
                frame = Image.blend(frame, nxt, ease((trans - (end - t)) / trans))
            return frame.convert("RGB")
    return gif_bg().convert("RGB")


def save_gif_outputs():
    if FRAME_DIR.exists():
        shutil.rmtree(FRAME_DIR)
    FRAME_DIR.mkdir(parents=True, exist_ok=True)
    fps_mp4, fps_gif, duration = 24, 12, 9.2
    writer = None
    try:
        import imageio.v2 as imageio

        writer = imageio.get_writer(MP4_PATH, fps=fps_mp4, codec="libx264", quality=8, macro_block_size=16)
    except Exception:
        pass
    gif_frames = []
    for i in range(int(duration * fps_mp4)):
        frame = render_gif_at(i / fps_mp4)
        if writer:
            writer.append_data(np.asarray(frame))
        if i % (fps_mp4 // fps_gif) == 0:
            frame_path = FRAME_DIR / f"{len(gif_frames) + 1:03d}.png"
            frame.save(frame_path, quality=95)
            gif_frames.append(frame)
    if writer:
        writer.close()
    pal = [f.convert("P", palette=Image.ADAPTIVE, colors=128) for f in gif_frames]
    pal[0].save(GIF_PATH, save_all=True, append_images=pal[1:], duration=int(1000 / fps_gif), loop=0, optimize=True, disposal=2)
    return GIF_PATH, MP4_PATH if MP4_PATH.exists() else None, FRAME_DIR, len(gif_frames)


def main():
    PACKAGE.mkdir(parents=True, exist_ok=True)
    previews = save_deck_previews()
    contact = create_contact_sheet(previews)
    build_ppt()
    notes = save_notes_blueprint()
    gif, mp4, frames, frame_count = save_gif_outputs()
    with ZipFile(PPTX_PATH) as z:
        slides = [n for n in z.namelist() if n.startswith("ppt/slides/slide") and n.endswith(".xml")]
        text_runs = sum(z.read(n).decode("utf-8", errors="ignore").count("<a:t>") for n in slides)
    print(PPTX_PATH)
    print(DECK_DIR)
    print(contact)
    print(notes)
    print(gif)
    print(mp4)
    print(frames)
    print(f"slides={len(slides)}")
    print(f"text_runs={text_runs}")
    print(f"gif_frames={frame_count}")


if __name__ == "__main__":
    main()
