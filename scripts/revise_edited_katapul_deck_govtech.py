from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor


PACKAGE = Path("product-showcase-package")
SOURCE = PACKAGE / "KataPult_Modest_Product_Walkthrough_7slides_EDITED.pptx"
OUT = PACKAGE / "KataPult_Modest_Product_Walkthrough_EDITED_GovTech_Refined.pptx"
SCREEN_DIR = PACKAGE / "govtech_screenshots"


INK = RGBColor(24, 32, 43)
MUTED = RGBColor(83, 95, 115)
LINE = RGBColor(216, 211, 200)
PAPER = RGBColor(248, 247, 244)
GREEN = RGBColor(127, 174, 45)


def delete_shape(shape):
    el = shape._element
    el.getparent().remove(el)


def set_text(shape, text):
    if not hasattr(shape, "text_frame"):
        return
    tf = shape.text_frame
    if not tf.paragraphs:
        return
    p = tf.paragraphs[0]
    if not p.runs:
        p.text = text
    else:
        p.runs[0].text = text
        for run in p.runs[1:]:
            run.text = ""


def add_text(slide, x, y, w, h, text, size, color=INK, bold=False, align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    for attr in ["margin_left", "margin_right", "margin_top", "margin_bottom"]:
        setattr(tf, attr, 0)
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.space_after = Pt(0)
    p.line_spacing = 1.05
    if align is not None:
        p.alignment = align
    for run in p.runs:
        run.font.name = "Segoe UI"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return box


def add_browser(slide, x, y, w, h, image_path):
    shadow = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x + 0.06), Inches(y + 0.08), Inches(w), Inches(h))
    shadow.fill.solid()
    shadow.fill.fore_color.rgb = RGBColor(232, 228, 220)
    shadow.line.fill.background()

    frame = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    frame.fill.solid()
    frame.fill.fore_color.rgb = RGBColor(255, 255, 255)
    frame.line.color.rgb = LINE
    frame.line.width = Pt(0.75)

    bar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.34))
    bar.fill.solid()
    bar.fill.fore_color.rgb = RGBColor(245, 243, 238)
    bar.line.fill.background()

    for i, c in enumerate([RGBColor(239, 68, 68), RGBColor(245, 158, 11), RGBColor(34, 197, 94)]):
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.18 + i * 0.2), Inches(y + 0.12), Inches(0.1), Inches(0.1))
        dot.fill.solid()
        dot.fill.fore_color.rgb = c
        dot.line.fill.background()

    slide.shapes.add_picture(str(image_path), Inches(x + 0.18), Inches(y + 0.46), Inches(w - 0.36), Inches(h - 0.62))


def add_label(slide, x, y, w, text):
    pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.32))
    pill.fill.solid()
    pill.fill.fore_color.rgb = RGBColor(238, 246, 232)
    pill.line.color.rgb = RGBColor(198, 219, 180)
    pill.line.width = Pt(0.5)
    add_text(slide, x + 0.14, y + 0.095, w - 0.28, 0.12, text, 7.8, INK, True, PP_ALIGN.CENTER)


def main():
    prs = Presentation(SOURCE)

    # Slide 3: keep user's centralisation story, sharpen the feedback output label.
    s3 = prs.slides[2]
    for sh in s3.shapes:
        if hasattr(sh, "text") and sh.text.strip() == "feedback and writing-support pages.":
            set_text(sh, "answer feedback and writing-support pages.")
        elif hasattr(sh, "text") and sh.text.strip() == "Feedback explanations":
            set_text(sh, "Answer feedback and explanations")
        elif hasattr(sh, "text") and "routes them into drills, tests, feedback" in sh.text:
            set_text(sh, "KataPult centralises relevant learning materials in one place, then routes them into drills, tests, answer feedback and writing-support pages.")

    # Slide 4: replace the weak screenshot area with a large real product state showing full explanation.
    s4 = prs.slides[3]
    for sh in list(s4.shapes):
        # Remove old browser mockup and small decorative chips above the loop.
        if sh.left >= Inches(6.65) and sh.top < Inches(5.0):
            delete_shape(sh)
    for sh in s4.shapes:
        if hasattr(sh, "text") and "understand the answer" in sh.text:
            set_text(sh, "Learners can pick a module, practise briefly, see right-or-wrong feedback with an AI-assisted explanation, then return through daily challenge, XP or streaks.")
        elif hasattr(sh, "text") and sh.text.strip() == "Right/wrong + LLM explanation.":
            set_text(sh, "Feedback plus AI explanation.")
        # Move the loop cards lower so the screenshot can breathe.
        if sh.top >= Inches(5.0):
            sh.top = sh.top + Inches(0.82)

    add_browser(s4, 6.72, 0.98, 5.88, 4.92, SCREEN_DIR / "04-imbuhan-real-explanation-focused.png")
    add_label(s4, 6.9, 5.58, 5.52, "Actual Imbuhan answer state: feedback plus full explanation")
    add_text(s4, 12.56, 0.32, 0.18, 0.2, "04", 8.5, MUTED, True, PP_ALIGN.RIGHT)

    # Slide 5: keep the user's table, but replace the small decorative screenshot with a more useful essay recall view.
    s5 = prs.slides[4]
    for sh in list(s5.shapes):
        if sh.left >= Inches(8.2) and sh.top >= Inches(2.0):
            delete_shape(sh)
    add_browser(s5, 8.32, 2.08, 4.28, 4.24, SCREEN_DIR / "05-essay-recall-focused.png")
    add_label(s5, 8.5, 6.12, 3.9, "Essay recall packs show reusable phrases and examples")

    # Slide 6: keep user's current-state structure, make wording more natural and stakeholder-safe.
    s6 = prs.slides[5]
    for sh in s6.shapes:
        if hasattr(sh, "text") and "testing and taking users" in sh.text:
            set_text(sh, "Current state: trialling with users and collecting feedback")
        elif hasattr(sh, "text") and sh.text.strip() == "The current numbers show coverage and informal trial use.":
            set_text(sh, "The current numbers show content coverage and early informal use.")
        elif hasattr(sh, "text") and sh.text.strip() == "Review collected items and generated explanations before treating them as validated material.":
            set_text(sh, "Review collected items and AI-assisted explanations before treating them as validated material.")
        elif hasattr(sh, "text") and sh.text.strip() == "Use feedback and usage patterns to improve the most-used modules first.":
            set_text(sh, "Use learner feedback and usage patterns to improve the most-used modules first.")

    prs.save(OUT)
    print(OUT.resolve())


if __name__ == "__main__":
    main()
